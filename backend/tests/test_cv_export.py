"""— CV exports + deterministic ATS lint."""

import json

import pytest
from docx import Document as open_docx
from io import BytesIO


def test_custom_text_docx_runs_bullets_and_ats_stripping():
    from app.services.cv_export_service import to_ats_text, to_docx

    text = (
        "Curious **engineer** from Athens.\n"
        "- *Ships* fast\n"
        "- Loves [Python](https://python.org)\n"
    )
    payload = {
        "snapshot": {},
        "blocks": [
            {"kind": "custom_text", "props": {"title": "About me", "text": text}}
        ],
    }

    docx_bytes = to_docx(payload, "CV")
    document = open_docx(BytesIO(docx_bytes))
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    about = paragraphs[1]
    assert about.text == "Curious engineer from Athens."
    bold = [run for run in about.runs if run.bold]
    assert bold and bold[0].text == "engineer"
    bullets = [p for p in paragraphs if p.style.name == "List Bullet"]
    assert [p.text for p in bullets] == [
        "Ships fast",
        "Loves Python: https://python.org",
    ]
    italic = [run for p in bullets for run in p.runs if run.italic]
    assert italic and italic[0].text == "Ships"

    ats = to_ats_text(payload)
    assert "**" not in ats and "*" not in ats
    assert "Curious engineer" in ats
    assert "- Ships fast" in ats
    assert "Loves Python: https://python.org" in ats


def test_language_line_cefr_upgrades_to_certificate_band():
    from app.services.cv_export_service import _language_line

    lang = {"label": "English", "code": "en", "level": "advanced", "cefr": "C1"}
    ecpe = {
        "title": (
            "Examination for the Certificate of Proficiency in English (ECPE) - C2"
        ),
        "org": "University of Michigan",
        "start": "May 2025",
        "end": "",
    }
    snapshot = {"certifications": [ecpe]}

    assert _language_line(lang, {}, snapshot) == "English — advanced"
    assert _language_line(lang, {"show_cefr": True}, snapshot) == (
        "English — advanced (C2)"
    )
    assert _language_line(
        lang, {"show_cefr": True, "show_proficiency": True}, snapshot
    ) == (
        "English — advanced (C2) · Examination for the Certificate of "
        "Proficiency in English (ECPE) - C2, May 2025"
    )
    assert (
        _language_line(
            {"label": "Greek", "code": "el", "level": "native"},
            {"show_cefr": True},
            snapshot,
        )
        == "Greek — native (Native)"
    )


async def _experience(client, headers, **overrides) -> dict:
    body = {
        "title": "DevOps intern",
        "kind": "internship",
        "org_name": "Acme Cloud",
        "start": "2025-01-01",
        "end": "2025-12-31",
        "hours_per_week": 40,
        "description": "Deployed things.",
        "skills": [],
        "achievements": [{"text": "Cut deploy time 40%"}],
        **overrides,
    }
    created = await client.post("/api/v1/me/experience", json=body, headers=headers)
    assert created.status_code == 201, created.text
    return created.json()


async def _make_built_cv(client, headers) -> dict:
    response = await client.put(
        "/api/v1/profile",
        json={
            "basics": {
                "birth_year": 2005,
                "education_level": "bachelor",
                "country": "Greece",
                "city": "Athens",
                "email": "me@example.com",
                "phone": "+30 555 0100",
                "headline": "Backend student",
            },
            "academics": {"languages": [{"code": "en", "level": "advanced"}]},
            "aspirations": [{"label": "Ship reliable systems", "notes": ""}],
            "interests": [],
            "hobbies": [],
            "likes": [],
            "dislikes": [],
            "work_preferences": {},
            "constraints": {},
        },
        headers=headers,
    )
    assert response.status_code == 200, response.text
    await _experience(client, headers)
    created = await client.post(
        "/api/v1/cv", json={"title": "Backend Intern CV"}, headers=headers
    )
    assert created.status_code == 201, created.text
    return created.json()


async def test_export_markdown_auto_versions(
    client, auth_headers, profile_ready, seeded_catalog
):
    cv = await _make_built_cv(client, auth_headers)
    response = await client.post(
        f"/api/v1/cv/{cv['id']}/export", json={"format": "md"}, headers=auth_headers
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/markdown")
    assert "backend-intern-cv.md" in response.headers["content-disposition"]
    body = response.text
    assert "# " in body and "DevOps intern" in body and "Acme Cloud" in body
    assert "Cut deploy time 40%" in body
    versions = await client.get(f"/api/v1/cv/{cv['id']}/versions", headers=auth_headers)
    rows = versions.json()
    assert len(rows) == 1 and rows[0]["created_by"] == "export"


async def test_export_all_formats(client, auth_headers, profile_ready, seeded_catalog):
    cv = await _make_built_cv(client, auth_headers)
    for fmt, media, marker in [
        ("ats_text", "text/plain", "DevOps intern"),
        (
            "docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            None,
        ),
        ("json", "application/json", "career_assistant.cv_export"),
    ]:
        response = await client.post(
            f"/api/v1/cv/{cv['id']}/export", json={"format": fmt}, headers=auth_headers
        )
        assert response.status_code == 200, (fmt, response.text)
        assert response.headers["content-type"].startswith(media), fmt
        if marker:
            assert marker in response.text, fmt
    docx_response = await client.post(
        f"/api/v1/cv/{cv['id']}/export", json={"format": "docx"}, headers=auth_headers
    )
    document = open_docx(BytesIO(docx_response.content))
    headings = [
        p.text for p in document.paragraphs if p.style.name.startswith("Heading")
    ]
    assert any("Experience" in heading for heading in headings)
    assert any("DevOps intern" in heading for heading in headings)

    json_response = await client.post(
        f"/api/v1/cv/{cv['id']}/export", json={"format": "json"}, headers=auth_headers
    )
    package = json.loads(json_response.content)
    assert package["content"]["snapshot"]["experience"][0]["org"] == "Acme Cloud"

    pdf_response = await client.post(
        f"/api/v1/cv/{cv['id']}/export", json={"format": "pdf"}, headers=auth_headers
    )
    if pdf_response.status_code == 503:
        pytest.skip(
            "optional server PDF engine not installed "
            "(pip install -r requirements-pdf.txt && playwright install chromium)"
        )
    assert pdf_response.headers["content-type"].startswith("application/pdf")
    assert pdf_response.content.startswith(b"%PDF")


async def test_export_unknown_format_rejected(
    client, auth_headers, profile_ready, seeded_catalog
):
    cv = await _make_built_cv(client, auth_headers)
    response = await client.post(
        f"/api/v1/cv/{cv['id']}/export", json={"format": "rtf"}, headers=auth_headers
    )
    assert response.status_code == 422


async def test_lint_flags_missing_contact_and_passes_filled(
    client, auth_headers, seeded_catalog
):
    anonymous = await client.post(
        "/api/v1/auth/register",
        json={
            "email": "anon@example.com",
            "password": "supersecret1",
            "full_name": "",
        },
    )
    assert anonymous.status_code == 201, anonymous.text
    bare_headers = {"Authorization": f"Bearer {anonymous.json()['access_token']}"}
    response = await client.put(
        "/api/v1/profile",
        json={
            "basics": {"birth_year": 2005, "education_level": "bachelor"},
            "academics": {},
            "aspirations": [],
            "interests": [],
            "hobbies": [],
            "likes": [],
            "dislikes": [],
            "work_preferences": {},
            "constraints": {},
        },
        headers=bare_headers,
    )
    assert response.status_code == 200, response.text
    bare = await client.post(
        "/api/v1/cv", json={"title": "Bare CV"}, headers=bare_headers
    )
    report = (
        await client.get(f"/api/v1/cv/{bare.json()['id']}/lint", headers=bare_headers)
    ).json()
    assert report["passed"] is False and report["score"] < 90
    ids = {check["id"] for check in report["checks"]}
    assert "contact_name" in ids and "contact_phone" in ids

    built = await _make_built_cv(client, auth_headers)
    report = (
        await client.get(f"/api/v1/cv/{built['id']}/lint", headers=auth_headers)
    ).json()
    assert report["passed"] is True, report["checks"]
    assert report["score"] >= 90
    order = next(check for check in report["checks"] if check["id"] == "section_order")
    assert order["level"] == "pass"


async def test_lint_flags_non_standard_heading(
    client, auth_headers, profile_ready, seeded_catalog
):
    cv = await _make_built_cv(client, auth_headers)
    await client.patch(
        f"/api/v1/cv/{cv['id']}",
        json={
            "working_content": {
                "blocks": [
                    {"kind": "header"},
                    {
                        "kind": "custom_text",
                        "props": {"title": "My Journey", "text": "Hello"},
                    },
                    {
                        "kind": "items",
                        "props": {"title": "Experience", "source_key": "experience"},
                    },
                ],
                "overrides": {},
            }
        },
        headers=auth_headers,
    )
    report = (
        await client.get(f"/api/v1/cv/{cv['id']}/lint", headers=auth_headers)
    ).json()
    heading_checks = [
        check for check in report["checks"] if check["id"] == "standard_headings"
    ]
    assert heading_checks and "my journey" in heading_checks[0]["message"]


async def test_pdf_export_downloads_real_pdf_bytes(
    client, auth_headers, profile_ready, seeded_catalog, monkeypatch
):
    import app.services.cv_export_service as export_module

    async def _fake_pdf(html: str) -> bytes:
        return b"%PDF-1.4 fake-bytes"

    monkeypatch.setattr(export_module, "html_to_pdf", _fake_pdf)
    cv = await _make_built_cv(client, auth_headers)
    response = await client.post(
        f"/api/v1/cv/{cv['id']}/export", json={"format": "pdf"}, headers=auth_headers
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/pdf")
    assert "attachment" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith('.pdf"')
    assert response.content.startswith(b"%PDF")


async def test_pdf_export_503_when_engine_unavailable(
    client, auth_headers, profile_ready, seeded_catalog, monkeypatch
):
    import app.services.cv_export_service as export_module
    from app.services.cv_pdf_service import PDFEngineUnavailable

    async def _unavailable(html: str) -> bytes:
        raise PDFEngineUnavailable("Server PDF engine not installed — use print")

    monkeypatch.setattr(export_module, "html_to_pdf", _unavailable)
    cv = await _make_built_cv(client, auth_headers)
    response = await client.post(
        f"/api/v1/cv/{cv['id']}/export", json={"format": "pdf"}, headers=auth_headers
    )
    assert response.status_code == 503, response.text
    assert "print" in response.json()["detail"].lower()
    assert "attachment" not in response.headers.get("content-disposition", "")


def test_lint_section_order_is_area_aware():
    """A sidebar layout renders the sidebar column first (side=left) —
    the section-order check compares against that reading order, not
    the stored block order (the AI redesigns assign sidebar areas)."""
    from types import SimpleNamespace

    from app.services.cv_export_service import lint

    blocks = [
        {"kind": "header", "props": {}},
        {"kind": "summary", "props": {"title": "Summary"}},  # main, stored first
        {"kind": "skills", "props": {"title": "Skills"}, "area": "sidebar"},
    ]
    snapshot = {
        "basics": {"name": "Test", "email": "t@example.com"},
        "summary": "x",
        "skills": [{"label": "Python", "level": 5}],
    }
    template = SimpleNamespace(
        content={
            "blocks": blocks,
            "design": {"layout": "sidebar", "sidebar_side": "left"},
            "pages": {"default_max_pages": 1, "overflow_policy": "warn"},
        }
    )
    html = "<h2>Skills</h2><h2>Summary</h2>"  # sidebar column renders first
    report = lint({"snapshot": snapshot, "blocks": blocks}, html, {}, template)
    order = next(check for check in report["checks"] if check["id"] == "section_order")
    assert order["level"] == "pass", report["checks"]

    right = SimpleNamespace(
        content={
            "blocks": blocks,
            "design": {"layout": "sidebar", "sidebar_side": "right"},
            "pages": {"default_max_pages": 1, "overflow_policy": "warn"},
        }
    )
    html_right = "<h2>Summary</h2><h2>Skills</h2>"
    report = lint({"snapshot": snapshot, "blocks": blocks}, html_right, {}, right)
    order = next(check for check in report["checks"] if check["id"] == "section_order")
    assert order["level"] == "pass", report["checks"]


def test_hidden_blocks_skip_exports_and_keep_lint_order_pass():
    from dataclasses import asdict

    from app.schemas.cv_template import TemplateContent
    from app.services.cv_export_service import lint, to_ats_text, to_docx, to_markdown
    from app.services.cv_renderer import render_cv

    payload = {
        "snapshot": {},
        "blocks": [
            {
                "kind": "custom_text",
                "props": {"title": "About me", "text": "Visible **text**."},
            },
            {
                "kind": "custom_text",
                "props": {"title": "Secret", "text": "Hidden text."},
                "hidden": True,
            },
        ],
    }

    markdown = to_markdown(payload)
    assert "About me" in markdown and "Visible **text**." in markdown
    assert "Secret" not in markdown and "Hidden text." not in markdown

    ats = to_ats_text(payload)
    assert "Secret" not in ats and "Hidden text." not in ats

    document = open_docx(BytesIO(to_docx(payload, "CV")))
    assert "Secret" not in "\n".join(p.text for p in document.paragraphs)

    content = TemplateContent.model_validate({"blocks": payload["blocks"]})
    rendered = render_cv(content, {})
    report = lint(payload, rendered.html, asdict(rendered.metrics), None)
    by_id = {check["id"]: check for check in report["checks"]}
    assert by_id["section_order"]["level"] == "pass", report["checks"]
    assert all("Secret" not in check["message"] for check in report["checks"])
